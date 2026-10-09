import csv
import io
import json

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.auth import get_current_user, get_user_org
from app.database import get_db
from app.models.entities import Answer, AnswerStatus, Question, Questionnaire, User
from app.schemas.api import ApproveAnswerRequest, QuestionOut, QuestionnaireOut
from app.services.answer_builder import build_draft_for_question
from app.services.question_matcher import suggest_control_id

router = APIRouter(prefix="/questionnaires", tags=["questionnaires"])


def _parse_questions_from_csv(content: bytes) -> list[str]:
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    questions: list[str] = []
    for row in reader:
        if not row:
            continue
        cell = row[0].strip()
        if cell and cell.lower() not in ("question", "questions", "security question"):
            questions.append(cell)
    return questions


def _parse_questions_from_xlsx(content: bytes) -> list[str]:
    wb = load_workbook(filename=io.BytesIO(content), read_only=True, data_only=True)
    sheet = wb.active
    questions: list[str] = []
    for row in sheet.iter_rows(min_row=1, max_col=1, values_only=True):
        val = row[0]
        if val is None:
            continue
        cell = str(val).strip()
        if cell and cell.lower() not in ("question", "questions", "security question"):
            questions.append(cell)
    return questions


@router.get("", response_model=list[QuestionnaireOut])
def list_questionnaires(user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    org = get_user_org(db, user)
    items = db.query(Questionnaire).filter(Questionnaire.organization_id == org.id).all()
    return [
        QuestionnaireOut(
            id=q.id,
            title=q.title,
            buyer_name=q.buyer_name,
            uploaded_at=q.uploaded_at,
            question_count=len(q.questions),
        )
        for q in items
    ]


@router.post("/upload", response_model=QuestionnaireOut)
async def upload_questionnaire(
    title: str = Form(...),
    buyer_name: str = Form(""),
    file: UploadFile = File(...),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    org = get_user_org(db, user)
    raw = await file.read()
    name = (file.filename or "").lower()
    if name.endswith(".xlsx"):
        questions_text = _parse_questions_from_xlsx(raw)
    elif name.endswith(".csv"):
        questions_text = _parse_questions_from_csv(raw)
    else:
        raise HTTPException(status_code=400, detail="Upload .csv or .xlsx with questions in column A")

    if not questions_text:
        raise HTTPException(status_code=400, detail="No questions found in file")

    questionnaire = Questionnaire(organization_id=org.id, title=title, buyer_name=buyer_name)
    db.add(questionnaire)
    db.flush()

    for idx, qtext in enumerate(questions_text):
        q = Question(
            questionnaire_id=questionnaire.id,
            row_index=idx + 1,
            question_text=qtext,
            suggested_control_id=suggest_control_id(qtext),
        )
        db.add(q)
    db.commit()
    db.refresh(questionnaire)

    for q in questionnaire.questions:
        build_draft_for_question(db, org.id, q)

    return QuestionnaireOut(
        id=questionnaire.id,
        title=questionnaire.title,
        buyer_name=questionnaire.buyer_name,
        uploaded_at=questionnaire.uploaded_at,
        question_count=len(questionnaire.questions),
    )


@router.get("/{questionnaire_id}/questions", response_model=list[QuestionOut])
def get_questions(
    questionnaire_id: int,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    org = get_user_org(db, user)
    questionnaire = (
        db.query(Questionnaire)
        .filter(Questionnaire.id == questionnaire_id, Questionnaire.organization_id == org.id)
        .first()
    )
    if not questionnaire:
        raise HTTPException(status_code=404, detail="Questionnaire not found")
    out: list[QuestionOut] = []
    for q in sorted(questionnaire.questions, key=lambda x: x.row_index):
        ans = q.answer
        out.append(
            QuestionOut(
                id=q.id,
                row_index=q.row_index,
                question_text=q.question_text,
                suggested_control_id=q.suggested_control_id,
                answer_status=ans.status.value if ans else None,
                draft_text=ans.draft_text if ans else None,
                approved_text=ans.approved_text if ans else None,
            )
        )
    return out


@router.post("/{questionnaire_id}/regenerate")
def regenerate_answers(
    questionnaire_id: int,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    org = get_user_org(db, user)
    questionnaire = (
        db.query(Questionnaire)
        .filter(Questionnaire.id == questionnaire_id, Questionnaire.organization_id == org.id)
        .first()
    )
    if not questionnaire:
        raise HTTPException(status_code=404, detail="Questionnaire not found")
    count = 0
    for q in questionnaire.questions:
        build_draft_for_question(db, org.id, q)
        count += 1
    return {"regenerated": count}


@router.post("/questions/{question_id}/approve")
def approve_answer(
    question_id: int,
    body: ApproveAnswerRequest,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    org = get_user_org(db, user)
    question = (
        db.query(Question)
        .join(Questionnaire)
        .filter(Question.id == question_id, Questionnaire.organization_id == org.id)
        .first()
    )
    if not question or not question.answer:
        raise HTTPException(status_code=404, detail="Question not found")
    answer = question.answer
    answer.approved_text = body.approved_text or answer.draft_text
    answer.status = AnswerStatus.approved
    db.commit()
    return {"status": "approved"}


@router.get("/{questionnaire_id}/export.csv")
def export_csv(
    questionnaire_id: int,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    org = get_user_org(db, user)
    questionnaire = (
        db.query(Questionnaire)
        .filter(Questionnaire.id == questionnaire_id, Questionnaire.organization_id == org.id)
        .first()
    )
    if not questionnaire:
        raise HTTPException(status_code=404, detail="Questionnaire not found")

    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["row", "question", "control_id", "status", "answer", "evidence"])
    for q in sorted(questionnaire.questions, key=lambda x: x.row_index):
        ans: Answer | None = q.answer
        final = (ans.approved_text if ans and ans.status == AnswerStatus.approved else None) or (
            ans.draft_text if ans else ""
        )
        writer.writerow(
            [
                q.row_index,
                q.question_text,
                q.suggested_control_id or "",
                ans.status.value if ans else "",
                final,
                ans.evidence_refs if ans else "[]",
            ]
        )
    buffer.seek(0)
    filename = f"dealroom-export-{questionnaire_id}.csv"
    return StreamingResponse(
        iter([buffer.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
